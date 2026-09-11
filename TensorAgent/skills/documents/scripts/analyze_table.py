#!/usr/bin/env python3
"""Read a CSV or XLSX, compute over its columns, and write the answer back out.

    analyze_table.py sales.csv
    analyze_table.py sales.csv --group-by Region --sum Revenue --mean Units
    analyze_table.py book.xlsx --sheet Q3 --group-by Region --sum Revenue --out summary.xlsx
    analyze_table.py sales.csv --where "Region=West" --sum Revenue --out summary.csv

With no --group-by it profiles every column: type, how many values are filled,
and for numeric columns count/sum/mean/median/min/max/stdev. With --group-by it
aggregates the requested columns per distinct value of that column.

The output format follows the --out extension (.json, .csv or .xlsx); with no
--out the JSON goes to stdout, which is what a model should read.

Reading .xlsx uses openpyxl with data_only=True, which returns the value cached
in the file. A workbook written by a tool that does not cache values -- openpyxl
on its own does not -- has None in every formula cell; this says so instead of
treating it as empty.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import statistics
import sys

NUMERIC_STRIP = " \t$%,"


def _read_csv(path: str, delimiter: "str | None", encoding: str) -> "tuple[list, list]":
    with open(path, newline="", encoding=encoding) as handle:
        sample = handle.read(64 * 1024)
        handle.seek(0)
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = ","
        reader = csv.reader(handle, delimiter=delimiter)
        rows = [row for row in reader]
    return (rows[0], rows[1:]) if rows else ([], [])


def read_csv(path: str, delimiter: "str | None") -> "tuple[list, list]":
    """Read a UTF-8 CSV, falling back to the common Windows export encoding.

    A BOM is authoritative for UTF-16/32, while ``utf-8-sig`` retains the
    existing behavior of removing a UTF-8 BOM from the first heading. UTF-8
    stays strict so a legacy byte is never silently replaced. Election systems
    and spreadsheet applications also commonly export Windows-1252; retrying
    from the beginning is necessary because the first non-UTF-8 byte may occur
    after the delimiter sample.
    """
    with open(path, "rb") as handle:
        prefix = handle.read(4)
    if prefix.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        return _read_csv(path, delimiter, "utf-32")
    if prefix.startswith((b"\xff\xfe", b"\xfe\xff")):
        return _read_csv(path, delimiter, "utf-16")
    try:
        return _read_csv(path, delimiter, "utf-8-sig")
    except UnicodeDecodeError:
        return _read_csv(path, delimiter, "cp1252")


def read_xlsx(path: str, sheet_name: "str | None") -> "tuple[list, list, int]":
    """Header, body rows, and how many formula cells hold no cached value.

    That third number is the difference between a real zero and a hole. openpyxl
    with data_only=True returns whatever the last application to save the file
    cached, and openpyxl on its own caches nothing — so a workbook another tool
    generated can have a formula in every total and None behind all of them.
    Summing that silently would report zero and look like an answer.
    """
    from openpyxl import load_workbook

    values = load_workbook(path, data_only=True, read_only=True)
    formulas = load_workbook(path, data_only=False, read_only=True)
    try:
        if sheet_name:
            if sheet_name not in values.sheetnames:
                raise SystemExit(f"analyze_table: {path} has no sheet named {sheet_name!r}; "
                                 f"it has {', '.join(values.sheetnames)}")
            value_sheet, formula_sheet = values[sheet_name], formulas[sheet_name]
        else:
            value_sheet, formula_sheet = values.worksheets[0], formulas.worksheets[0]
        # One pass over both views at once: a read-only sheet is a stream, and
        # walking it twice is both slower and not something to rely on.
        rows: list = []
        uncached = 0
        for value_row, formula_row in zip(
            value_sheet.iter_rows(values_only=True), formula_sheet.iter_rows(values_only=True)
        ):
            rows.append(list(value_row))
            for cached, formula in zip(value_row, formula_row):
                if cached is None and isinstance(formula, str) and formula.startswith("="):
                    uncached += 1
    finally:
        values.close()
        formulas.close()

    rows = [row for row in rows if any(value is not None and str(value).strip() for value in row)]
    if not rows:
        return [], [], uncached
    header = ["" if value is None else str(value) for value in rows[0]]
    return header, rows[1:], uncached


def as_number(value: object) -> "float | None":
    """A number if the cell holds one, else None.

    Currency and thousands separators are stripped because a CSV exported from a
    spreadsheet is full of them and refusing to read `$1,240.00` as a number
    would make this script useless on exactly the files it is for.
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return None if (isinstance(value, float) and math.isnan(value)) else float(value)
    text = str(value).strip()
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    for char in NUMERIC_STRIP:
        text = text.replace(char, "")
    try:
        number = float(text)
    except ValueError:
        return None
    return -number if negative else number


def column_values(rows: list, index: int) -> list:
    return [row[index] if index < len(row) else None for row in rows]


def profile(header: list, rows: list) -> list:
    out = []
    for index, name in enumerate(header):
        values = column_values(rows, index)
        filled = [v for v in values if v is not None and str(v).strip() != ""]
        numbers = [n for n in (as_number(v) for v in filled) if n is not None]
        entry: dict = {
            "column": name,
            "index": index,
            "filled": len(filled),
            "empty": len(values) - len(filled),
            "distinct": len({str(v) for v in filled}),
            "numeric": len(numbers) == len(filled) and bool(filled),
        }
        if numbers:
            entry["stats"] = {
                "count": len(numbers),
                "sum": sum(numbers),
                "mean": statistics.fmean(numbers),
                "median": statistics.median(numbers),
                "min": min(numbers),
                "max": max(numbers),
                "stdev": statistics.stdev(numbers) if len(numbers) > 1 else 0.0,
            }
        else:
            counts: dict = {}
            for value in filled:
                counts[str(value)] = counts.get(str(value), 0) + 1
            entry["top_values"] = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))[:5]
        out.append(entry)
    return out


def resolve(header: list, name: str) -> int:
    if name in header:
        return header.index(name)
    lowered = [h.strip().lower() for h in header]
    if name.strip().lower() in lowered:
        return lowered.index(name.strip().lower())
    raise SystemExit(f"analyze_table: no column named {name!r}; the file has {', '.join(header)}")


ORDERING = {">", "<", ">=", "<="}


def apply_filters(header: list, rows: list, clauses: list) -> "tuple[list, int]":
    """Keep rows matching every `Column=value` / `Column>value` clause.

    Returns the kept rows and how many were dropped because a cell an ordering
    clause needed was blank or not a number. That count is reported rather than
    swallowed: "8 of 10 rows" and "8 of 10 rows, 1 unreadable" are different
    answers, and the second one is often the finding.
    """
    if not clauses:
        return rows, 0
    tests = []
    for clause in clauses:
        for operator in (">=", "<=", "!=", "=", ">", "<"):
            if operator in clause:
                name, _, wanted = clause.partition(operator)
                wanted = wanted.strip()
                if operator in ORDERING and as_number(wanted) is None:
                    raise SystemExit(f"analyze_table: --where {clause!r} compares with {operator}, "
                                     f"so {wanted!r} has to be a number")
                tests.append((resolve(header, name.strip()), operator, wanted))
                break
        else:
            raise SystemExit(f"analyze_table: --where {clause!r} needs one of = != > < >= <=")
    kept: list = []
    incomparable = 0
    for row in rows:
        dropped_for_type = False
        for index, operator, wanted in tests:
            value = row[index] if index < len(row) else None
            number, target = as_number(value), as_number(wanted)
            if operator in ORDERING:
                if number is None:
                    dropped_for_type = True
                    break
                ok = {">": number > target, "<": number < target,
                      ">=": number >= target, "<=": number <= target}[operator]
            elif number is not None and target is not None:
                ok = (number == target) if operator == "=" else (number != target)
            else:
                text = "" if value is None else str(value).strip()
                ok = (text.lower() == wanted.lower()) if operator == "=" else (text.lower() != wanted.lower())
            if not ok:
                break
        else:
            kept.append(row)
            continue
        if dropped_for_type:
            incomparable += 1
    return kept, incomparable


def group(header: list, rows: list, by: str, sums: list, means: list) -> dict:
    key_index = resolve(header, by)
    sum_indexes = [(name, resolve(header, name)) for name in sums]
    mean_indexes = [(name, resolve(header, name)) for name in means]

    buckets: dict = {}
    order: list = []
    for row in rows:
        key = "" if key_index >= len(row) or row[key_index] is None else str(row[key_index]).strip()
        if key not in buckets:
            buckets[key] = {"rows": 0, "values": {}}
            order.append(key)
        bucket = buckets[key]
        bucket["rows"] += 1
        for name, index in sum_indexes + mean_indexes:
            number = as_number(row[index] if index < len(row) else None)
            if number is not None:
                bucket["values"].setdefault(name, []).append(number)

    columns = [by, "rows"] + [f"sum({n})" for n, _ in sum_indexes] + [f"mean({n})" for n, _ in mean_indexes]
    out_rows = []
    for key in sorted(order):
        bucket = buckets[key]
        row = [key, bucket["rows"]]
        for name, _ in sum_indexes:
            row.append(sum(bucket["values"].get(name, [])))
        for name, _ in mean_indexes:
            values = bucket["values"].get(name, [])
            row.append(statistics.fmean(values) if values else None)
        out_rows.append(row)

    totals = ["TOTAL", len(rows)]
    for name, index in sum_indexes:
        numbers = [n for n in (as_number(r[index] if index < len(r) else None) for r in rows) if n is not None]
        totals.append(sum(numbers))
    for name, index in mean_indexes:
        numbers = [n for n in (as_number(r[index] if index < len(r) else None) for r in rows) if n is not None]
        totals.append(statistics.fmean(numbers) if numbers else None)

    return {"columns": columns, "rows": out_rows, "totals": totals, "groups": len(order)}


def write_output(path: str, result: dict) -> None:
    extension = os.path.splitext(path)[1].lower()
    grouped = result.get("summary")
    if extension == ".json":
        with open(path, "w", encoding="utf-8") as handle:
            json.dump(result, handle, indent=2, default=str)
        return
    if grouped is None:
        raise SystemExit(f"analyze_table: --out {path} needs --group-by; a column profile only writes as .json")
    if extension == ".csv":
        with open(path, "w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(grouped["columns"])
            writer.writerows(grouped["rows"])
            writer.writerow(grouped["totals"])
        return
    if extension == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Summary"
        sheet.append(grouped["columns"])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="1F3864")
            cell.fill = PatternFill("solid", fgColor="DEEAF6")
            cell.alignment = Alignment(horizontal="center")
        for row in grouped["rows"]:
            sheet.append(row)
        sheet.append(grouped["totals"])
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
        for index, name in enumerate(grouped["columns"], start=1):
            widest = max([len(str(name))] + [len(str(r[index - 1])) for r in grouped["rows"]] or [0])
            sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = min(max(widest + 3, 10), 40)
        sheet.freeze_panes = "A2"
        workbook.save(path)
        return
    raise SystemExit(f"analyze_table: --out {path} has an extension this does not write "
                     "(.json, .csv and .xlsx are the ones it does)")


def main() -> int:
    parser = argparse.ArgumentParser(description="Compute over a CSV or XLSX table.")
    parser.add_argument("input", help="a .csv or .xlsx file")
    parser.add_argument("--sheet", help="sheet name, for an .xlsx (default: the first)")
    parser.add_argument("--delimiter", help="CSV delimiter (default: sniffed)")
    parser.add_argument("--where", action="append", default=[],
                        help="keep rows matching Column=value (also != > < >= <=); repeatable")
    parser.add_argument("--group-by", help="column to group by")
    parser.add_argument("--sum", action="append", default=[], help="column to total per group; repeatable")
    parser.add_argument("--mean", action="append", default=[], help="column to average per group; repeatable")
    parser.add_argument("--out", help="write the result to this .json, .csv or .xlsx")
    args = parser.parse_args()

    if not os.path.isfile(args.input):
        raise SystemExit(f"analyze_table: {args.input} does not exist")
    extension = os.path.splitext(args.input)[1].lower()
    uncached = 0
    if extension in (".xlsx", ".xlsm"):
        header, rows, uncached = read_xlsx(args.input, args.sheet)
        source = "xlsx"
    elif extension in (".csv", ".tsv", ".txt"):
        header, rows = read_csv(args.input, args.delimiter or ("\t" if extension == ".tsv" else None))
        source = "csv"
    else:
        raise SystemExit(f"analyze_table: {args.input} is neither a .csv nor an .xlsx; "
                         "convert it first or name the right file")

    if not header:
        raise SystemExit(f"analyze_table: {args.input} is empty")

    filtered, incomparable = apply_filters(header, rows, args.where)
    result: dict = {
        "file": os.path.abspath(args.input),
        "source": source,
        "sheet": args.sheet,
        "columns": header,
        "rows_read": len(rows),
        "rows_used": len(filtered),
        "rows_dropped_not_a_number": incomparable,
        "formula_cells_without_cached_values": uncached,
    }
    if uncached:
        print(f"analyze_table: {uncached} cell(s) in {os.path.basename(args.input)} hold a formula and no "
              "cached value, so they read as blank here and count as nothing. Whatever wrote the file did "
              "not store its results; open it in a spreadsheet application once, or rebuild it with "
              "make_xlsx.py, which does cache them.", file=sys.stderr)
    if args.group_by:
        # Under its own key: "columns" at the top level is the file's header, and
        # overwriting it with the summary's columns would lose it.
        result["summary"] = group(header, filtered, args.group_by, args.sum, args.mean)
    else:
        if args.sum or args.mean:
            raise SystemExit("analyze_table: --sum and --mean need --group-by; "
                             "without it the column profile already reports totals and means")
        result["profile"] = profile(header, filtered)

    if args.out:
        write_output(args.out, result)
        result["written"] = os.path.abspath(args.out)
    print(json.dumps(result, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
